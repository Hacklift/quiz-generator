from .export_labels import export_labels, export_type
from io import BytesIO
from typing import Any

from docx import Document


def generate_docx(payload: dict[str, Any]):
    buffer = BytesIO()
    doc = Document()

    labels = export_labels(payload)
    title = payload.get("title") or labels["title"]
    quiz_type = export_type(payload)
    description = payload.get("description")

    doc.add_heading(title, level=1)
    doc.add_paragraph(f"{labels['type']}: {quiz_type}")
    if description:
        doc.add_paragraph(description)
    doc.add_heading(labels["questions"], level=2)

    for item in payload.get("questions", []):
        question_number = item.get("number")
        question_prefix = f"{labels['question']} {question_number}" if question_number else labels["question"]

        doc.add_paragraph(f"{question_prefix}: {item['question']}")
        for option in item.get("options") or []:
            doc.add_paragraph(option)

        doc.add_paragraph(f"{labels['answer']}: {item['answer']}")
        if item.get("explanation"):
            doc.add_paragraph(f"{labels['explanation']}: {item['explanation']}")
        doc.add_paragraph("")

    doc.save(buffer)
    buffer.seek(0)
    return buffer
